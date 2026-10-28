from pathlib import Path

from docx import Document
from docx.enum.section import WD_SECTION
from docx.enum.table import WD_ALIGN_VERTICAL
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.oxml import OxmlElement
from docx.oxml.ns import qn
from docx.shared import Cm, Pt, RGBColor


ROOT = Path(__file__).resolve().parent
OUTPUT = ROOT / "Lab_6_Quartz_Report_Toy_Store.docx"


def set_cell_shading(cell, fill):
    properties = cell._tc.get_or_add_tcPr()
    shading = OxmlElement("w:shd")
    shading.set(qn("w:fill"), fill)
    properties.append(shading)


def set_cell_border(cell, color="D9D9D9"):
    properties = cell._tc.get_or_add_tcPr()
    borders = properties.first_child_found_in("w:tcBorders")
    if borders is None:
        borders = OxmlElement("w:tcBorders")
        properties.append(borders)
    for edge_name in ("top", "left", "bottom", "right", "insideH", "insideV"):
        edge = borders.find(qn(f"w:{edge_name}"))
        if edge is None:
            edge = OxmlElement(f"w:{edge_name}")
            borders.append(edge)
        edge.set(qn("w:val"), "single")
        edge.set(qn("w:sz"), "4")
        edge.set(qn("w:color"), color)


def set_font(run, bold=False, size=12):
    run.font.name = "Times New Roman"
    run._element.rPr.rFonts.set(qn("w:ascii"), "Times New Roman")
    run._element.rPr.rFonts.set(qn("w:hAnsi"), "Times New Roman")
    run._element.rPr.rFonts.set(qn("w:cs"), "Times New Roman")
    run.font.size = Pt(size)
    run.font.bold = bold
    run.font.color.rgb = RGBColor(0, 0, 0)


def add_paragraph(document, text="", bold_prefix=None, align=None, size=12, first_indent=True):
    paragraph = document.add_paragraph()
    paragraph.paragraph_format.space_after = Pt(6)
    paragraph.paragraph_format.line_spacing = 1.15
    if first_indent:
        paragraph.paragraph_format.first_line_indent = Cm(1.0)
    if align is not None:
        paragraph.alignment = align
    if bold_prefix and text.startswith(bold_prefix):
        prefix = paragraph.add_run(bold_prefix)
        set_font(prefix, bold=True, size=size)
        rest = paragraph.add_run(text[len(bold_prefix):])
        set_font(rest, size=size)
    else:
        run = paragraph.add_run(text)
        set_font(run, size=size)
    return paragraph


def add_heading(document, text, level=1):
    paragraph = document.add_paragraph(style=f"Heading {level}")
    paragraph.paragraph_format.space_before = Pt(12 if level == 1 else 8)
    paragraph.paragraph_format.space_after = Pt(6)
    paragraph.paragraph_format.keep_with_next = True
    run = paragraph.add_run(text)
    set_font(run, bold=True, size=15 if level == 1 else 13)
    return paragraph


def add_bullets(document, items):
    for item in items:
        paragraph = document.add_paragraph(style="List Bullet")
        paragraph.paragraph_format.space_after = Pt(3)
        paragraph.paragraph_format.line_spacing = 1.1
        run = paragraph.add_run(item)
        set_font(run, size=11)


def add_table(document, headers, rows, widths=None):
    table = document.add_table(rows=1, cols=len(headers))
    table.style = "Table Grid"
    table.autofit = False
    header_cells = table.rows[0].cells
    for index, header in enumerate(headers):
        cell = header_cells[index]
        cell.vertical_alignment = WD_ALIGN_VERTICAL.CENTER
        set_cell_shading(cell, "1F4E78")
        set_cell_border(cell)
        paragraph = cell.paragraphs[0]
        paragraph.alignment = WD_ALIGN_PARAGRAPH.CENTER
        run = paragraph.add_run(header)
        set_font(run, bold=True, size=10)
        run.font.color.rgb = RGBColor(255, 255, 255)
        if widths:
            cell.width = Cm(widths[index])
    for row_index, row in enumerate(rows):
        cells = table.add_row().cells
        for index, value in enumerate(row):
            cell = cells[index]
            cell.vertical_alignment = WD_ALIGN_VERTICAL.CENTER
            set_cell_border(cell)
            if row_index % 2 == 1:
                set_cell_shading(cell, "F2F6FA")
            paragraph = cell.paragraphs[0]
            paragraph.paragraph_format.space_after = Pt(2)
            paragraph.paragraph_format.line_spacing = 1.0
            run = paragraph.add_run(str(value))
            set_font(run, size=9)
            if widths:
                cell.width = Cm(widths[index])
    document.add_paragraph().paragraph_format.space_after = Pt(2)
    return table


def add_code(document, lines):
    paragraph = document.add_paragraph()
    paragraph.paragraph_format.left_indent = Cm(0.6)
    paragraph.paragraph_format.right_indent = Cm(0.4)
    paragraph.paragraph_format.space_after = Pt(6)
    paragraph.paragraph_format.line_spacing = 1.0
    properties = paragraph._p.get_or_add_pPr()
    shading = OxmlElement("w:shd")
    shading.set(qn("w:fill"), "F2F2F2")
    properties.append(shading)
    run = paragraph.add_run(lines)
    run.font.name = "Courier New"
    run._element.rPr.rFonts.set(qn("w:ascii"), "Courier New")
    run.font.size = Pt(9)
    return paragraph


def add_page_number(paragraph):
    paragraph.alignment = WD_ALIGN_PARAGRAPH.CENTER
    run = paragraph.add_run("Сторінка ")
    set_font(run, size=9)
    field = OxmlElement("w:fldSimple")
    field.set(qn("w:instr"), "PAGE")
    paragraph._p.append(field)


def configure_document(document):
    section = document.sections[0]
    section.top_margin = Cm(2)
    section.bottom_margin = Cm(2)
    section.left_margin = Cm(2.5)
    section.right_margin = Cm(1.5)
    styles = document.styles
    for name in ("Normal", "Heading 1", "Heading 2", "Title"):
        style = styles[name]
        style.font.name = "Times New Roman"
        style._element.rPr.rFonts.set(qn("w:ascii"), "Times New Roman")
        style._element.rPr.rFonts.set(qn("w:hAnsi"), "Times New Roman")
        style.font.color.rgb = RGBColor(0, 0, 0)
    footer = section.footer
    add_page_number(footer.paragraphs[0])


def build_report():
    document = Document()
    configure_document(document)

    cover = document.add_paragraph()
    cover.alignment = WD_ALIGN_PARAGRAPH.CENTER
    cover.paragraph_format.space_after = Pt(14)
    run = cover.add_run("ЗВІТ ДО ЛАБОРАТОРНОЇ РОБОТИ №6")
    set_font(run, bold=True, size=16)
    title = document.add_paragraph(style="Title")
    title.alignment = WD_ALIGN_PARAGRAPH.CENTER
    title.paragraph_format.space_after = Pt(20)
    run = title.add_run("Quartz Scheduler у проєкті Toy Store")
    set_font(run, bold=True, size=20)
    add_paragraph(document, "Тема: планування, збереження та надійне виконання регулярних завдань", align=WD_ALIGN_PARAGRAPH.CENTER, size=13, first_indent=False)
    document.add_paragraph().paragraph_format.space_after = Pt(60)
    add_paragraph(document, "Проєкт: Toy Store", size=13, first_indent=False)
    add_paragraph(document, "Технології: Java 17, Spring Boot 3, Quartz 2.5, MariaDB, Liquibase, JUnit 5, Mockito", size=13, first_indent=False)
    document.add_paragraph().paragraph_format.space_after = Pt(85)
    add_paragraph(document, "Виконав(ла): ______________________________", size=12, first_indent=False)
    add_paragraph(document, "Група: ____________________________________", size=12, first_indent=False)
    add_paragraph(document, "2026", align=WD_ALIGN_PARAGRAPH.CENTER, size=12, first_indent=False)
    document.add_page_break()

    add_heading(document, "Зміст")
    contents = [
        "1. Мета та завдання роботи",
        "2. Вихідні дані та обґрунтування регулярних задач",
        "3. Архітектура рішення",
        "4. Налаштування Quartz JDBC Job Store і Liquibase",
        "5. Реалізовані jobs і triggers",
        "6. Надійність, ідемпотентність і обробка помилок",
        "7. Тестування та результати перевірки",
        "8. Послідовність демонстрації викладачу",
        "9. Висновок",
    ]
    add_bullets(document, contents)

    add_heading(document, "1. Мета та завдання роботи")
    add_paragraph(document, "Метою лабораторної роботи є інтеграція Quartz Scheduler у застосунок Toy Store для запуску регулярних доменних задач за розкладом, який зберігається у базі даних. Рішення має бути стійким до перезапуску застосунку, не створювати дублікати робіт, мати визначену поведінку у разі пропущеного запуску та не допускати паралельної обробки однакових даних.")
    add_paragraph(document, "У межах роботи виконано налаштування Quartz JDBC Job Store для MariaDB, підключено керування схемою через Liquibase, створено два класи Quartz Job, три triggers, передачу параметрів через JobDataMap та тестове покриття всіх сценаріїв, визначених у завданні.")

    add_heading(document, "2. Вихідні дані та обґрунтування регулярних задач")
    add_paragraph(document, "Toy Store уже має сутності замовлення та іграшки. Для планування було обрано дві задачі, які відповідають доменній логіці магазину та можуть виконуватися без участі користувача.")
    add_table(document,
              ["Задача", "Причина автоматизації", "Результат"],
              [
                  ["Очищення застарілих замовлень", "Замовлення зі статусом PLACED не повинні залишатися активними необмежений час.", "Такі замовлення переводяться у DECLINED після заданого часу очікування."],
                  ["Звіт про малий залишок", "Менеджеру потрібно регулярно бачити іграшки, запас яких наближається до вичерпання.", "У лог записується кількість позицій у звіті за заданим порогом залишку."],
              ], [4.1, 6.2, 6.2])
    add_paragraph(document, "Завдання очищення змінює лише записи зі статусом PLACED. Після зміни статусу повторний запуск не знаходить цей самий запис, тому операція є ідемпотентною. Звітова задача працює у режимі лише читання і не модифікує дані.")

    add_heading(document, "3. Архітектура рішення")
    add_paragraph(document, "Конфігурація Quartz зосереджена в одному класі QuartzConfiguration. У ньому створюються JobDetail і Trigger beans з фіксованими назвою та групою. Стабільні ключі важливі для JDBC Job Store: після перезапуску Quartz відновлює існуючий опис із таблиць, а не створює нові дублікати.")
    add_table(document,
              ["Компонент", "Призначення", "Відповідальність"],
              [
                  ["QuartzConfiguration", "Опис jobs і triggers", "Визначає JobKey, TriggerKey, JobDataMap, misfire та timezone."],
                  ["AutowiringSpringBeanJobFactory", "Інтеграція Quartz зі Spring", "Створює екземпляр job і виконує autowiring сервісів Spring."],
                  ["StaleOrderCleanupJob", "Quartz Job", "Зчитує параметри і викликає один метод OrderMaintenanceService."],
                  ["LowStockReportJob", "Quartz Job", "Зчитує параметри і викликає один метод InventoryReportService."],
                  ["OrderMaintenanceService", "Доменна логіка", "У транзакції переводить застарілі PLACED замовлення у DECLINED."],
                  ["InventoryReportService", "Доменна логіка", "У read-only транзакції формує дані про малий залишок."],
              ], [4.7, 4.5, 7.3])
    add_paragraph(document, "Jobs навмисно залишені мінімалістичними: вони не містять JPA-логіки, не працюють напряму з репозиторіями й не відкривають транзакції. Їхня роль - валідувати параметри JobDataMap і делегувати роботу одному сервісу Spring.")

    add_heading(document, "4. Налаштування Quartz JDBC Job Store і Liquibase")
    add_paragraph(document, "До Maven-проєкту додано залежність spring-boot-starter-quartz. У application.yml обрано тип сховища jdbc та вимкнено автоматичне створення схеми Quartz засобами Spring Boot. Таким чином, єдине джерело керування структурою бази даних - Liquibase.")
    add_code(document, "spring:\n  quartz:\n    job-store-type: jdbc\n    overwrite-existing-jobs: false\n    jdbc:\n      initialize-schema: never")
    add_paragraph(document, "До master changelog додано changeset 007. Він запускає SQL-схему Quartz 2.5.2 для MariaDB InnoDB. Схема створює таблиці QRTZ_JOB_DETAILS, QRTZ_TRIGGERS, QRTZ_SIMPLE_TRIGGERS, QRTZ_CRON_TRIGGERS, QRTZ_FIRED_TRIGGERS, QRTZ_LOCKS та інші службові таблиці й індекси. Саме в них зберігаються jobs, triggers, дані запусків і блокування.")
    add_paragraph(document, "Встановлено scheduler name ToyStoreScheduler та JDBC delegate StdJDBCDelegate. Параметр overwrite-existing-jobs: false не дозволяє застосунку перезаписувати вже збережені визначення без явного рішення, що зменшує ризик дублювання або непомітної зміни розкладу після деплою.")

    add_heading(document, "5. Реалізовані jobs і triggers")
    add_table(document,
              ["Job", "JobDataMap за замовчуванням", "Trigger", "Розклад і misfire"],
              [
                  ["stale-order-cleanup", "staleHours=24, batchSize=100", "stale-order-cleanup-hourly", "Щогодини; NEXT_WITH_EXISTING_COUNT; batchSize=25."],
                  ["stale-order-cleanup", "staleHours=24, batchSize=100", "stale-order-cleanup-nightly", "Щодня о 02:00 Europe/Kyiv; DO_NOTHING; staleHours=48."],
                  ["low-stock-report", "minimumStock=5, reportName=low-stock-toys", "low-stock-report-daily", "Щодня о 18:00 Europe/Kyiv; DO_NOTHING; minimumStock=3."],
              ], [3.5, 4.5, 4.4, 5.0])
    add_paragraph(document, "Перший Job має два triggers, як вимагає завдання. Це демонструє, що один тип доменної операції може запускатися за різними правилами. Для нічного запуску визначено більший поріг давності замовлення, а для погодинного - менший розмір пакета.")
    add_paragraph(document, "Quartz об'єднує JobDataMap JobDetail і JobDataMap Trigger. Значення тригера мають пріоритет. Наприклад, погодинний trigger замінює batchSize=100 на batchSize=25, але зберігає staleHours=24 із JobDetail.")
    add_code(document, "JobDetail:  staleHours=24, batchSize=100\nTrigger:    batchSize=25\nMerged map: staleHours=24, batchSize=25")
    add_paragraph(document, "Для cron-triggers явно встановлено TimeZone Europe/Kyiv. Це виключає залежність від локального часового поясу сервера. Для кожного trigger також задана інструкція misfire: cron-завдання не наздоганяють пропущений запуск, а очікують наступного планового часу; погодинне simple-завдання продовжує свій графік без лавиноподібного виконання пропущених задач.")

    add_heading(document, "6. Надійність, ідемпотентність і обробка помилок")
    add_paragraph(document, "Обидва класи jobs позначено @DisallowConcurrentExecution. Якщо новий момент запуску настає, коли попередній запуск того самого job ще триває, Quartz не запускає другий екземпляр паралельно. Це важливо для очищення замовлень, щоб один запис не оброблявся двічі одночасно.")
    add_paragraph(document, "Для OrderMaintenanceService транзакційна межа розташована на сервісному методі declineStalePlacedOrders. Метод вибирає лише замовлення PLACED, старші за заданий поріг, обмежує їх кількістю batchSize, змінює статус на DECLINED і зберігає результат. Після цього повторний запуск не знаходить ті самі записи. Отже, другий запуск не здійснює додаткової зміни даних.")
    add_paragraph(document, "Якщо сервіс кидає RuntimeException, job фіксує помилку у лог і передає Quartz JobExecutionException із refireImmediately=false. Обрана стратегія - не створювати нескінченний цикл повторних спроб. Quartz завершує поточний запуск і переходить до наступного запланованого trigger-а.")
    add_paragraph(document, "Відсутні або некоректні параметри JobDataMap також обробляються явно. Допоміжний клас QuartzJobData перевіряє наявність ключів, тип і додатність числових значень. Якщо ключ відсутній або має неправильний формат, job завершується зрозумілим JobExecutionException.")

    add_heading(document, "7. Тестування та результати перевірки")
    add_paragraph(document, "Для лабораторної додано модульні тести конфігурації, Quartz jobs, сервісу обслуговування замовлень і конфігурації персистентності. Перевірка виконувалася командами mvn test та mvn verify. Обидві команди завершилися успішно; verify додатково перевірив встановлений поріг покриття JaCoCo.")
    add_table(document,
              ["Сценарій із завдання", "Де перевіряється", "Очікуваний результат"],
              [
                  ["Jobs і triggers в одному місці", "QuartzConfigurationTest", "JobKey і TriggerKey мають визначені назву та групу; trigger посилається на свій JobKey."],
                  ["Пріоритет параметрів trigger", "QuartzConfigurationTest", "batchSize з trigger замінює default, staleHours з JobDetail зберігається."],
                  ["Легкі jobs", "QuartzJobsTest", "Кожен job зчитує параметри та викликає один сервіс."],
                  ["Відсутній параметр", "QuartzJobsTest", "Викидається JobExecutionException із назвою відсутнього ключа."],
                  ["Заборона паралельності", "QuartzJobsTest", "Перевіряється анотація @DisallowConcurrentExecution на обох jobs."],
                  ["Misfire і timezone", "QuartzConfigurationTest", "Перевіряються інструкції та Europe/Kyiv у cron-triggers."],
                  ["Ідемпотентність", "OrderMaintenanceServiceTest", "Другий запуск для тих самих даних повертає 0 і не змінює запис повторно."],
                  ["Збій job", "QuartzJobsTest", "JobExecutionException має refireImmediately=false."],
                  ["JDBC schema через Liquibase", "QuartzPersistenceConfigurationTest", "Перевіряється SQL QRTZ_* і initialize-schema: never."],
              ], [4.7, 4.5, 8.2])

    add_heading(document, "8. Послідовність демонстрації викладачу")
    add_bullets(document, [
        "Показати залежність spring-boot-starter-quartz у pom.xml та блок spring.quartz у application.yml.",
        "Відкрити changelog-master.xml і changeset 007, пояснити, що Liquibase створює QRTZ_* таблиці, а Spring Boot не ініціалізує схему самостійно.",
        "Відкрити QuartzConfiguration: показати два JobDetail, три Trigger beans, їхні ключі та JobDataMap.",
        "Показати, що stale-order-cleanup має два triggers: погодинний і нічний.",
        "Пояснити злиття JobDataMap: trigger може перевизначити параметр job за замовчуванням.",
        "Відкрити StaleOrderCleanupJob і LowStockReportJob: показати @DisallowConcurrentExecution, перевірку параметрів і один виклик сервісу.",
        "Відкрити OrderMaintenanceService: показати @Transactional та ідемпотентність через вибірку лише PLACED замовлень.",
        "Запустити mvn test або показати результати тестів QuartzConfigurationTest, QuartzJobsTest, OrderMaintenanceServiceTest.",
    ])

    add_heading(document, "9. Висновок")
    add_paragraph(document, "У Toy Store реалізовано надійне планування регулярних задач за допомогою Quartz Scheduler. Розклад і службовий стан зберігаються у MariaDB через JDBC Job Store, а схема керується Liquibase. Дві доменні задачі мають стабільні ключі, параметризацію через JobDataMap, явні правила misfire та часовий пояс Europe/Kyiv. Логіка винесена до транзакційних сервісів, jobs захищені від паралельного запуску, повторне очищення є ідемпотентним, а помилки не створюють нескінченних повторних запусків. Реалізацію підтверджено автоматичними тестами та Maven verify.")

    document.save(OUTPUT)
    print(OUTPUT)


if __name__ == "__main__":
    build_report()
